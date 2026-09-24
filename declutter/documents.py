"""PDF / DOCX / XLSX text extraction, content-based duplicate detection and
keyword classification."""

import hashlib
import io
import os
import re

from . import config as C
from .grouping import BKTree, UnionFind, jaccard, normalize_words, shingles, simhash
from .photos import _sha256

# ------------------------------------------------------------ extraction ----


def _pdf_text(path, use_ocr):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("password-protected PDF")
    parts, pages = [], reader.pages
    for page in list(pages)[: C.DOC_MAX_PDF_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
        if sum(len(p) for p in parts) > C.DOC_MAX_CHARS:
            break
    text = "\n".join(parts)
    note = ""
    if len(text.strip()) < 20 and use_ocr:  # probably a scan: OCR embedded page images
        text2 = _ocr_pdf_images(list(pages)[: C.DOC_OCR_SCANNED_PAGES])
        if text2.strip():
            text, note = text2, "scanned PDF (OCR)"
    return text, note


def _ocr_pdf_images(pages):
    from PIL import Image

    from .resources import configure_tesseract

    cmd, _ = configure_tesseract()
    if not cmd:
        return ""
    import pytesseract

    out = []
    for page in pages:
        try:
            images = page.images
        except Exception:
            continue
        for img in list(images)[:3]:
            try:
                im = Image.open(io.BytesIO(img.data))
                im.thumbnail((2000, 2000))
                out.append(pytesseract.image_to_string(im.convert("L"), timeout=60))
            except Exception:
                continue
    return "\n".join(out)


def _docx_text(path):
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text for c in row.cells))
    for s in d.sections:
        try:
            parts += [p.text for p in s.header.paragraphs]
        except Exception:
            pass
    return "\n".join(parts)


def _xlsx_text(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts, n = [], 0
    try:
        for ws in wb.worksheets:
            parts.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    parts.append(" ".join(cells))
                n += len(row)
                if n > C.DOC_MAX_XLSX_CELLS:
                    break
    finally:
        wb.close()
    return "\n".join(parts)


def extract_text(path, use_ocr=True):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf_text(path, use_ocr)
    if ext == ".docx":
        return _docx_text(path), ""
    if ext == ".xlsx":
        return _xlsx_text(path), ""
    return "", ""


# -------------------------------------------------------- classification ----
# (pattern, weight). Patterns are regexes matched against lower-case text.
RULES = {
    "invoice": [
        (r"\binvoice\b", 3), (r"\btax invoice\b", 2), (r"\bbill (to|no|number|date)\b", 2), (r"\breceipt\b", 2),
        (r"\b(sub ?total|grand total|total amount|amount due|balance due)\b", 2), (r"\bgst(in)?\b|\bvat\b|\bcgst\b|\bsgst\b", 2),
        (r"\bqty\b|\bquantity\b", 1), (r"\bpayment (due|terms|received)\b", 2), (r"\bdue date\b", 1),
        (r"(₹|rs\.?|inr|\$|usd|€|£)\s?\d", 1), (r"\bpurchase order\b|\bp\.?o\.? number\b", 1),
        (r"\bbilling\b|\bbilled\b", 1), (r"\bstatement\b", 1),
    ],
    "id": [
        (r"\bcertificate\b", 3), (r"\bcertif(y|ied|ies)\b", 2), (r"\bthis is to certify\b", 3),
        (r"\bawarded\b|\bhas successfully completed\b|\bin recognition\b", 2), (r"\bpassport\b", 3),
        (r"\baadhaa?r\b|\bpan card\b|\bpermanent account number\b", 3), (r"\bdriving licen[cs]e\b|\bdriver'?s license\b", 3),
        (r"\bdate of birth\b|\bd\.?o\.?b\b", 2), (r"\bgovernment of\b|\brepublic of\b", 2), (r"\bidentity\b|\bid (card|no|number)\b", 2),
        (r"\bnationality\b|\bplace of birth\b", 2), (r"\bdiploma\b|\bdegree\b|\bmarksheet\b|\btranscript\b", 2),
        (r"\bvoter\b|\bsocial security\b|\bresidence permit\b|\bvisa\b", 2),
    ],
    "resume": [
        (r"\bresume\b|\brésumé\b|\bcurriculum vitae\b|\bcv\b", 3), (r"\b(work|professional) experience\b|\bexperience\b", 2),
        (r"\beducation\b", 2), (r"\bskills\b", 2), (r"\bobjective\b|\bprofessional summary\b|\bcareer summary\b", 2),
        (r"\blinkedin\b", 2), (r"\breferences\b", 1), (r"\bprojects\b", 1), (r"\bcertifications\b|\bachievements\b", 1),
        (r"\bemployment history\b|\binternship\b", 2), (r"\b(19|20)\d\d\s?[-–]\s?((19|20)\d\d|present)\b", 1),
    ],
    "notes": [
        (r"\bnotes?\b", 2), (r"\bmeeting\b|\bminutes\b", 2), (r"\bagenda\b", 2), (r"\bto-?do\b|\baction items?\b", 2),
        (r"\blecture\b|\bchapter\b|\bsyllabus\b|\bclass\b", 1), (r"\bsummary\b", 1), (r"\bideas?\b|\bdraft\b", 1),
        (r"\bremember\b|\breminder\b", 1), (r"\bjournal\b|\bdiary\b", 2),
    ],
}
NAME_HINTS = {
    "invoice": r"invoice|bill|receipt|inv[_\-\s]?\d|statement|payment",
    "id": r"certificate|cert\b|passport|aadhaa?r|pan\b|licen[cs]e|\bid\b|marksheet|diploma|degree",
    "resume": r"resume|résumé|\bcv\b|cv[_\-\s]|curriculum",
    "notes": r"notes?|meeting|minutes|lecture|todo|journal",
}
_RULES_C = {k: [(re.compile(p, re.I), w) for p, w in v] for k, v in RULES.items()}
_NAME_C = {k: re.compile(v, re.I) for k, v in NAME_HINTS.items()}


def classify(text, name):
    """Returns (category, confidence 0..1, reason)."""
    low = text.lower()[:60000]
    stem = os.path.splitext(name)[0].replace("_", " ")
    scores, hits = {}, {}
    for cat, rules in _RULES_C.items():
        s, h = 0, []
        for rx, w in rules:
            m = rx.search(low)
            if m:
                s += w
                h.append(m.group(0).strip())
        if _NAME_C[cat].search(stem):
            s += 3
            h.append("file name")
        scores[cat], hits[cat] = s, h
    best = max(scores, key=scores.get)
    ranked = sorted(scores.values(), reverse=True)
    top, second = ranked[0], ranked[1]
    if top < 4:
        return "other", 0.3, "no strong keywords"
    conf = min(0.95, 0.45 + 0.05 * top + 0.05 * (top - second))
    return best, round(conf, 2), "matched: " + ", ".join(dict.fromkeys(hits[best]))


# ------------------------------------------------------------ main entry ----
def analyze_document(ff, use_ocr=True):
    rec = {
        "path": ff.path, "root": ff.root, "size": ff.size, "mtime": ff.mtime,
        "name": os.path.basename(ff.path), "kind": "document", "error": None,
    }
    try:
        rec["sha256"] = _sha256(ff.path)
    except OSError as e:
        rec.update(error=str(e), category="other", confidence=0.0, reason="unreadable", words=0)
        return rec
    try:
        text, note = extract_text(ff.path, use_ocr)
    except Exception as e:
        text, note = "", ""
        rec["error"] = f"{type(e).__name__}: {e}"[:300]
    words = normalize_words(text[: C.DOC_MAX_CHARS])
    rec["words"] = len(words)
    if len(words) >= C.DOC_MIN_WORDS_FOR_TEXT_DUP:
        rec["text_hash"] = hashlib.sha1(" ".join(words).encode("utf-8")).hexdigest()
        k = 5 if len(words) >= 40 else 3
        sh = shingles(words, k)
        rec["_shingles"] = sh
        rec["simhash"] = simhash(sh)
    cat, conf, why = classify(text, rec["name"]) if (words or rec["name"]) else ("other", 0.2, "")
    if note:
        why = f"{note}; {why}"
    rec.update(category=cat, confidence=conf, reason=why)
    return rec


def _keeper_key(r):
    # the most complete (most words), then most recently modified, then largest
    return (r.get("words", 0), r["mtime"], r["size"])


def group_duplicate_docs(records):
    """Content-based duplicate grouping (in place). A re-saved copy with new
    metadata has a different file hash but the same extracted text."""
    uf = UnionFind(len(records))
    buckets = {}
    for i, r in enumerate(records):
        if r.get("sha256"):
            buckets.setdefault(("b", r["sha256"]), []).append(i)
        if r.get("text_hash"):
            buckets.setdefault(("t", r["text_hash"]), []).append(i)
    for members in buckets.values():
        for m in members[1:]:
            uf.union(members[0], m)

    tree = BKTree()
    text_idx = [i for i, r in enumerate(records) if r.get("simhash") is not None]
    for i in text_idx:
        tree.add(records[i]["simhash"], i)
    for i in text_idx:
        for j in tree.query(records[i]["simhash"], C.DOC_SIMHASH_MAX_DISTANCE):
            if j > i and uf.find(i) != uf.find(j):
                a, b = records[i], records[j]
                la, lb = a["words"], b["words"]
                if min(la, lb) / max(la, lb) < 0.7:
                    continue
                if jaccard(a["_shingles"], b["_shingles"]) >= C.DOC_JACCARD_MIN:
                    uf.union(i, j)

    n = 0
    for r in records:
        r["dup_group"] = None
        r["is_keeper"] = False
        r.pop("_shingles", None)
    for g in sorted(uf.groups(), key=lambda g: min(g)):
        n += 1
        members = [records[i] for i in g]
        keeper = max(members, key=_keeper_key)
        exact = len({m["sha256"] for m in members}) == 1
        for m in members:
            m["dup_group"] = n
            m["dup_exact"] = exact
            m["is_keeper"] = m is keeper
            if m is not keeper:
                m["category_detected"] = m["category"]
                m["category"] = keeper["category"]
    return n
