"""Generate a realistic two-folder test dataset ("old phone" vs "new phone").

Needs the dev extras (scikit-image for sample photos, reportlab for PDFs):
    pip install -r requirements-dev.txt
    python tests/make_fixtures.py /tmp/declutter_fixtures
"""

import os
import shutil
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def font(size, bold=False):
    cands = FONT_CANDIDATES[1:] + FONT_CANDIDATES[:1] if bold else FONT_CANDIDATES
    for p in cands:
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def sample(name):
    import skimage.data as d
    return Image.fromarray(getattr(d, name)()).convert("RGB")


def with_exif(img, path, make="Canon", model="EOS 80D", **kw):
    exif = Image.Exif()
    exif[271], exif[272] = make, model
    img.save(path, "JPEG", exif=exif.tobytes(), **kw)


def scenery(w=1200, h=800, seed=1):
    rng = np.random.default_rng(seed)
    a = np.zeros((h, w, 3), np.uint8)
    for y in range(h // 2):  # sky gradient
        t = y / (h / 2)
        a[y] = (int(90 + 60 * t), int(150 + 50 * t), 235)
    grass = np.zeros((h - h // 2, w, 3), np.float32)
    grass[..., 0], grass[..., 1], grass[..., 2] = 50, 140, 40
    grass += rng.normal(0, 25, grass.shape)
    a[h // 2:] = np.clip(grass, 0, 255).astype(np.uint8)
    img = Image.fromarray(a)
    d = ImageDraw.Draw(img)
    d.polygon([(0, h // 2), (300, h // 4), (650, h // 2)], fill=(90, 100, 110))
    d.ellipse((950, 60, 1050, 160), fill=(255, 240, 180))
    return img


def receipt():
    img = Image.new("RGB", (900, 1600), (250, 250, 245))
    d = ImageDraw.Draw(img)
    f, fb = font(34), font(44, True)
    y = 60
    lines = [("FRESH MART SUPERMARKET", fb), ("12 Market Road, Pune", f), ("Date: 14/03/2024  Time: 18:42", f),
             ("Invoice No: 55821", f), ("-" * 38, f)]
    items = [("Milk 1L", "56.00"), ("Bread", "40.00"), ("Eggs x12", "84.00"), ("Rice 5kg", "420.00"),
             ("Apples 1kg", "180.00"), ("Coffee 200g", "310.00"), ("Soap x3", "99.00")]
    for t, fnt in lines:
        d.text((60, y), t, fill=(20, 20, 20), font=fnt)
        y += 62
    for n, p in items:
        d.text((60, y), n, fill=(20, 20, 20), font=f)
        d.text((640, y), "Rs " + p, fill=(20, 20, 20), font=f)
        y += 58
    for t in ["-" * 38, "Subtotal          Rs 1189.00", "GST 5%            Rs 59.45",
              "TOTAL             Rs 1248.45", "Paid by card   Balance 0.00", "Thank you! Visit again"]:
        d.text((60, y), t, fill=(20, 20, 20), font=fb if "TOTAL" in t else f)
        y += 62
    return img.rotate(1.5, expand=True, fillcolor=(120, 110, 100))


def note_photo():
    img = Image.new("RGB", (1400, 1000), (245, 243, 230))
    d = ImageDraw.Draw(img)
    f = font(36)
    text = ("Lecture 7 - Thermodynamics. The first law states that energy cannot be created or "
            "destroyed, only transformed. Internal energy change equals heat added minus work done "
            "by the system. Remember to revise the examples on isothermal and adiabatic processes "
            "before Friday and read chapter nine on entropy and the second law of thermodynamics.")
    words, line, y = text.split(), "", 60
    for w in words:
        if d.textlength(line + " " + w, font=f) > 1260:
            d.text((70, y), line, fill=(30, 30, 90), font=f)
            y += 60
            line = w
        else:
            line = (line + " " + w).strip()
    d.text((70, y), line, fill=(30, 30, 90), font=f)
    return img


def screenshot():
    img = Image.new("RGB", (1080, 2340), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1080, 90), fill=(30, 30, 30))
    d.text((40, 25), "9:41", fill="white", font=font(40))
    d.rectangle((0, 90, 1080, 260), fill=(0, 120, 215))
    d.text((40, 140), "Messages", fill="white", font=font(60, True))
    for i in range(9):
        y = 320 + i * 220
        d.ellipse((40, y, 160, y + 120), fill=(200, 200, 210))
        d.text((190, y + 10), f"Contact {i + 1}", fill="black", font=font(44, True))
        d.text((190, y + 70), "See you tomorrow at the station!", fill=(90, 90, 90), font=font(36))
    return img


def meme(base):
    img = base.resize((800, 800))
    d = ImageDraw.Draw(img)
    f = font(84, True)
    for text, y in (("WHEN YOU FIND", 30), ("A DUPLICATE PHOTO", 680)):
        tw = d.textlength(text, font=f)
        x = (800 - tw) / 2
        for dx in (-4, 0, 4):
            for dy in (-4, 0, 4):
                d.text((x + dx, y + dy), text, fill="black", font=f)
        d.text((x, y), text, fill="white", font=f)
    return img


def make_pdf(path, lines, title="doc", author="someone"):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(title)
    c.setAuthor(author)
    y = 800
    for ln in lines:
        c.drawString(60, y, ln)
        y -= 18
    c.save()


def make_docx(path, paragraphs, author="someone"):
    import docx
    doc = docx.Document()
    doc.core_properties.author = author
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def make_xlsx(path, rows, creator="someone"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.properties.creator = creator
    wb.save(str(path))


INVOICE = ["TAX INVOICE", "Invoice Number: INV-2024-0042", "Invoice Date: 02 Feb 2024",
           "Bill To: Asha Traders, 4 MG Road, Bengaluru", "Description        Qty   Rate    Amount",
           "Web hosting (annual)  1   4800.00  4800.00", "Domain renewal        2    899.00  1798.00",
           "Subtotal: 6598.00", "GST 18%: 1187.64", "Total Amount Due: Rs 7785.64",
           "Payment due within 15 days. Bank: HDFC, A/c 00112233."]
RESUME = ["Priya Sharma", "Curriculum Vitae", "Email: priya@example.com | LinkedIn: linkedin.com/in/priya",
          "Objective: Data analyst seeking to apply statistics skills.",
          "Work Experience: Analyst, Acme Corp (2020 - 2024). Built dashboards and forecasting models.",
          "Education: B.Sc. Statistics, University of Pune, 2019.",
          "Skills: Python, SQL, Excel, Tableau, communication.", "Projects: Sales forecasting, churn analysis.",
          "References available on request."]
CERT = ["CERTIFICATE OF COMPLETION", "This is to certify that Rahul Verma",
        "has successfully completed the course Advanced Excel", "Awarded on 12 August 2023",
        "Certificate ID: CRT-88213", "Signature of Director, Skill Academy"]
NOTES = ["Meeting notes - project kickoff", "Agenda: scope, timeline, owners.",
         "Action items: Ravi to draft the plan by Monday; Meera to book the venue.",
         "Todo: share minutes with the team, schedule follow-up meeting next week.",
         "Summary: team agreed on a six week timeline with weekly check-ins."]


def build(out):
    out = Path(out)
    if out.exists():
        shutil.rmtree(out)
    old, new = out / "old_phone_backup", out / "new_phone_backup"
    for p in (old / "DCIM" / "Camera", old / "Screenshots", old / "Documents",
              new / "Pictures", new / "Documents" / "work", old / ".hidden"):
        p.mkdir(parents=True)

    astro, coffee, chelsea, rocket = sample("astronaut"), sample("coffee"), sample("chelsea"), sample("rocket")
    with_exif(astro, old / "DCIM/Camera/IMG_20230101_101010.jpg", quality=95)
    with_exif(astro.resize((400, 400)), new / "Pictures/astronaut_small_copy.jpg", quality=70)  # near-dup
    with_exif(coffee, old / "DCIM/Camera/IMG_20230205_083000.jpg", quality=92)
    shutil.copy(old / "DCIM/Camera/IMG_20230205_083000.jpg", new / "Pictures/coffee.jpg")  # exact dup
    with_exif(chelsea, old / "DCIM/Camera/IMG_20230310_120000.jpg", quality=92)
    with_exif(chelsea.filter(ImageFilter.GaussianBlur(6)), old / "DCIM/Camera/IMG_20230310_120001_blur.jpg", quality=92)
    with_exif(scenery(), old / "DCIM/Camera/IMG_20230420_170000.jpg", quality=92)
    scenery().convert("RGB").save(new / "Pictures/hike.webp", "WEBP", quality=80)  # near-dup, webp
    with_exif(rocket, new / "Pictures/PXL_20240101_launch.jpg", quality=90)
    screenshot().save(old / "Screenshots/Screenshot_20240115-094100.png")
    screenshot().save(new / "Pictures/unnamed_capture.png")  # dup by content, no screenshot name
    with_exif(receipt(), old / "DCIM/Camera/IMG_20240314_184500.jpg", make="Xiaomi", model="Redmi Note 9", quality=90)
    with_exif(note_photo(), new / "Pictures/IMG_20240320_090000.jpg", make="Apple", model="iPhone 12", quality=90)
    meme(chelsea).save(new / "Pictures/funny.jpg", quality=85)
    Image.new("RGB", (100, 80), (200, 30, 30)).save(new / "Pictures/tiny_icon.png")
    Image.new("RGB", (500, 500), (1, 2, 3)).save(old / ".hidden/ignored.png")  # hidden: skipped
    (old / "Documents/corrupt.jpg").write_bytes(b"this is not really a jpeg")

    make_pdf(old / "Documents/invoice_feb.pdf", INVOICE, title="Invoice", author="Billing v1")
    make_pdf(new / "Documents/work/invoice_feb (copy).pdf", INVOICE, title="Invoice copy", author="Re-saved v2")
    make_pdf(old / "Documents/certificate.pdf", CERT)
    receipt().save(old / "Documents/scanned_bill.pdf", "PDF", resolution=150)  # image-only PDF -> OCR
    make_docx(old / "Documents/Priya_Resume.docx", RESUME, author="Priya")
    make_docx(new / "Documents/work/CV_final_v2.docx", RESUME[:-1] + ["References available upon request."],
              author="Priya laptop")  # near-dup text
    make_docx(new / "Documents/work/meeting.docx", NOTES)
    rows = [["Month", "Rent", "Food", "Travel"], ["Jan", 15000, 6200, 1800], ["Feb", 15000, 5900, 2400],
            ["Mar", 15000, 6400, 1200], ["Apr", 15500, 6100, 2100]]
    make_xlsx(old / "Documents/budget.xlsx", rows, creator="A")
    make_xlsx(new / "Documents/budget_backup.xlsx", rows, creator="B")
    return out


if __name__ == "__main__":
    print(build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/declutter_fixtures"))
